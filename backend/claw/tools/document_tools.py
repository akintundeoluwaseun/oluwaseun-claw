"""Create and edit files: text, reports (Markdown or Word), spreadsheets (Excel or CSV).

Honesty rules enforced here, not left to the model:
- Report sources must be web pages or files Claw actually opened in this task.
- Spreadsheet values are checked against what Claw read; values that were not
  found are counted and reported so the user knows what to double-check.
- New files never overwrite existing ones. Editing an existing file first
  keeps a backup copy, and the change is read back to confirm it.
"""

from __future__ import annotations

import asyncio
import csv
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, Field, field_validator

from claw.agent.schemas import Capability
from claw.permissions.policy import PathNotAuthorized, RiskLevel
from claw.tools.base import ToolArgs, ToolContext, ToolError, ToolRegistry, ToolResult, ToolSpec
from claw.tools.file_tools import resolve_folder, safe_file_name, unique_path

TEXT_EXTS = {".txt", ".md", ".csv", ".json", ".html"}
Cell = str | int | float | bool | None


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def _record_output(ctx: ToolContext, path: Path, kind: str, verified: bool, note: str = "") -> None:
    ctx.memory.outputs.append({"path": str(path), "kind": kind, "verified": verified, "note": note})


def _new_path(ctx: ToolContext, folder: str | None, name: str, ext: str) -> Path:
    target_dir = resolve_folder(ctx, folder)
    path = unique_path(target_dir / safe_file_name(name, ext))
    ctx.policy.check_path(path, write=True)
    return path


# ---------- write_text_file ----------
class WriteTextArgs(ToolArgs):
    file_name: str = Field(min_length=1, max_length=100, description="e.g. summary.md")
    text: str = Field(min_length=1, max_length=500_000)
    folder: str | None = Field(default=None, max_length=500)

    @field_validator("file_name")
    @classmethod
    def _ext(cls, v: str) -> str:
        ext = Path(v).suffix.lower() or ".txt"
        if ext not in TEXT_EXTS:
            raise ValueError(f"text files must end in one of {', '.join(sorted(TEXT_EXTS))}")
        return v if Path(v).suffix else v + ".txt"


def check_write_text(args: WriteTextArgs, ctx: ToolContext) -> None:
    resolve_folder(ctx, args.folder)


async def write_text_file(args: WriteTextArgs, ctx: ToolContext) -> ToolResult:
    ext = Path(args.file_name).suffix.lower()
    path = _new_path(ctx, args.folder, args.file_name, ext)
    with open(path, "x", encoding="utf-8", newline="\n") as fh:
        fh.write(args.text)
    ok = path.read_text(encoding="utf-8") == args.text
    _record_output(ctx, path, "text", ok)
    return ToolResult(ok, f"Created {path.name} ({len(args.text):,} characters).", {"path": str(path)}, verified=ok)


# ---------- save_report ----------
class Section(BaseModel):
    heading: str = Field(min_length=1, max_length=200)
    body: str = Field(min_length=1, max_length=50_000)


class Source(BaseModel):
    title: str = Field(min_length=1, max_length=300)
    location: str = Field(min_length=3, max_length=2048, description="Web address or file path Claw opened")


class ReportArgs(ToolArgs):
    title: str = Field(min_length=3, max_length=200)
    sections: list[Section] = Field(min_length=1, max_length=40)
    sources: list[Source] = Field(default_factory=list, max_length=100)
    format: Literal["docx", "md"] = "docx"
    file_name: str | None = Field(default=None, max_length=100)
    folder: str | None = Field(default=None, max_length=500)


_URL = re.compile(r"^[a-z][a-z0-9+.\-]*://", re.I)


def _source_was_opened(location: str, ctx: ToolContext) -> bool:
    """A source may be cited only if Claw really opened it in this task.

    A file is matched on its canonical path, not the spelling the model used:
    on Windows the recorded path is case-folded, so a raw string match would
    reject a file Claw had just read.
    """
    if location in ctx.browser.seen_urls:
        return True
    if _URL.match(location):
        return False
    return ctx.memory.saw_file(location)


def check_report(args: ReportArgs, ctx: ToolContext) -> None:
    unseen = [s.location for s in args.sources if not _source_was_opened(s.location, ctx)]
    if unseen:
        raise ToolError("These sources were not opened during this task, so they can't be cited: " + ", ".join(unseen[:5]))
    resolve_folder(ctx, args.folder)


def _write_docx(path: Path, args: ReportArgs, stamp: str) -> None:
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    doc.core_properties.title = args.title
    doc.core_properties.author = "Oluwaseun Claw"
    doc.add_heading(args.title, level=0)
    note = doc.add_paragraph(f"Prepared by Oluwaseun Claw on {stamp}.")
    note.runs[0].font.size = Pt(9)
    for sec in args.sections:
        doc.add_heading(sec.heading, level=1)
        for para in [p for p in sec.body.split("\n\n") if p.strip()]:
            doc.add_paragraph(para.strip())
    if args.sources:
        doc.add_heading("Sources", level=1)
        for src in args.sources:
            doc.add_paragraph(f"{src.title} — {src.location}", style="List Number")
    with open(path, "xb") as fh:  # never overwrite
        doc.save(fh)


def _write_md(path: Path, args: ReportArgs, stamp: str) -> None:
    lines = [f"# {args.title}", "", f"_Prepared by Oluwaseun Claw on {stamp}._", ""]
    for sec in args.sections:
        lines += [f"## {sec.heading}", "", sec.body.strip(), ""]
    if args.sources:
        lines += ["## Sources", ""]
        for i, src in enumerate(args.sources, 1):
            loc = src.location if src.location.startswith("http") else f"`{src.location}`"
            lines.append(f"{i}. {src.title} — {loc}")
    with open(path, "x", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(lines) + "\n")


def _verify_report(path: Path, args: ReportArgs) -> bool:
    if path.suffix == ".docx":
        import docx

        text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    else:
        text = path.read_text(encoding="utf-8")
    return args.title in text and all(s.heading in text for s in args.sections) and all(s.location in text for s in args.sources)


async def save_report(args: ReportArgs, ctx: ToolContext) -> ToolResult:
    check_report(args, ctx)
    ext = "." + args.format
    path = _new_path(ctx, args.folder, args.file_name or args.title, ext)
    stamp = datetime.now().strftime("%d %B %Y, %H:%M")
    await asyncio.to_thread(_write_docx if args.format == "docx" else _write_md, path, args, stamp)
    ok = await asyncio.to_thread(_verify_report, path, args)
    note = "" if args.sources else "No sources were listed."
    _record_output(ctx, path, "report", ok, note)
    return ToolResult(
        ok,
        f"Saved the report '{args.title}' as {path.name} with {len(args.sections)} section(s) and {len(args.sources)} source(s).",
        {"path": str(path), "sources": len(args.sources)},
        verified=ok,
    )


# ---------- create_spreadsheet ----------
class SheetArgs(ToolArgs):
    file_name: str = Field(min_length=1, max_length=100)
    columns: list[str] = Field(min_length=1, max_length=50)
    rows: list[list[Cell]] = Field(min_length=1, max_length=5000)
    sheet_name: str = Field(default="Sheet1", min_length=1, max_length=31)
    format: Literal["xlsx", "csv"] = "xlsx"
    folder: str | None = Field(default=None, max_length=500)

    @field_validator("rows")
    @classmethod
    def _widths(cls, rows, info):
        cols = info.data.get("columns") or []
        for i, r in enumerate(rows):
            if len(r) > len(cols):
                raise ValueError(f"row {i + 1} has {len(r)} values but there are {len(cols)} columns")
        return rows

    @field_validator("sheet_name")
    @classmethod
    def _sheet(cls, v: str) -> str:
        if re.search(r"[\[\]:*?/\\]", v):
            raise ValueError("sheet names can't contain [ ] : * ? / \\")
        return v


def check_sheet(args: SheetArgs, ctx: ToolContext) -> None:
    resolve_folder(ctx, args.folder)


def provenance(rows: list[list[Cell]], corpus: str) -> tuple[int, int, list[str]]:
    """How many text values appear in what Claw read. Short/numeric values are not counted."""
    haystack = _norm(corpus)
    checked, found, missing = 0, 0, []
    for row in rows:
        for v in row:
            if not isinstance(v, str) or len(v.strip()) < 4:
                continue
            checked += 1
            if _norm(v) in haystack:
                found += 1
            elif len(missing) < 5:
                missing.append(v.strip()[:60])
    return checked, found, missing


def _safe_cell(v: Cell) -> Cell:
    # A cell starting with = + - @ would run as a formula in Excel: store it as plain text.
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _write_xlsx(path: Path, args: SheetArgs) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet_name
    ws.append(args.columns)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in args.rows:
        ws.append([_safe_cell(v) for v in row])
    ws.freeze_panes = "A2"
    for i, col in enumerate(args.columns, 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in args.rows[:200] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    with open(path, "xb") as fh:
        wb.save(fh)


def _write_csv(path: Path, args: SheetArgs) -> None:
    with open(path, "x", encoding="utf-8-sig", newline="") as fh:  # BOM so Excel opens UTF-8 correctly
        w = csv.writer(fh)
        w.writerow(args.columns)
        for row in args.rows:
            w.writerow(["" if v is None else _safe_cell(v) for v in row])


def _count_rows(path: Path) -> int:
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True)
        try:
            return wb.active.max_row - 1
        finally:
            wb.close()
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return sum(1 for _ in csv.reader(fh)) - 1


async def create_spreadsheet(args: SheetArgs, ctx: ToolContext) -> ToolResult:
    # A name that already says what it is wins over the format default. Asking for
    # "contacts.csv" used to give "contacts.csv.xlsx": the wrong kind of file, with
    # two extensions, and no word about it.
    named = Path(args.file_name).suffix.lower()
    fmt = named[1:] if named in (".csv", ".xlsx") else args.format
    path = _new_path(ctx, args.folder, args.file_name, "." + fmt)
    await asyncio.to_thread(_write_xlsx if fmt == "xlsx" else _write_csv, path, args)
    written = await asyncio.to_thread(_count_rows, path)
    ok = written == len(args.rows)
    checked, found, missing = provenance(args.rows, ctx.memory.corpus())
    note = ""
    if checked and found < checked:
        note = f"{checked - found} of {checked} text values were not found in anything Claw read (e.g. {', '.join(repr(m) for m in missing[:3])}). Check them."
    _record_output(ctx, path, "spreadsheet", ok, note)
    summary = f"Created {path.name} with {written} row(s) and {len(args.columns)} column(s)."
    if checked:
        summary += f" {found} of {checked} text values match what Claw read."
    return ToolResult(ok, summary, {"path": str(path), "rows": written, "checked": checked, "found": found, "missing": missing}, verified=ok)


# ---------- edit_text_file ----------
class EditArgs(ToolArgs):
    path: str = Field(min_length=1, max_length=500)
    find: str = Field(min_length=1, max_length=20_000)
    replace_with: str = Field(max_length=50_000)
    replace_all: bool = False


def _edit_target(ctx: ToolContext, raw: str) -> Path:
    try:
        path = ctx.policy.check_path(raw, write=True)
    except PathNotAuthorized as exc:
        raise ToolError(str(exc)) from exc
    if not path.is_file():
        raise ToolError(f"'{path}' is not a file.")
    if path.suffix.lower() not in TEXT_EXTS | {".log", ".ini", ".yaml", ".yml", ".toml", ".css", ".js"}:
        raise ToolError("Claw can only edit plain text files for now.")
    if path.stat().st_size > 5_000_000:
        raise ToolError("The file is too large to edit.")
    return path


def check_edit(args: EditArgs, ctx: ToolContext) -> None:
    path = _edit_target(ctx, args.path)
    count = path.read_text(encoding="utf-8-sig").count(args.find)
    if count == 0:
        raise ToolError("The text to replace was not found in the file.")
    if count > 1 and not args.replace_all:
        raise ToolError(f"The text appears {count} times. Give more surrounding text, or replace all.")


async def edit_text_file(args: EditArgs, ctx: ToolContext) -> ToolResult:
    check_edit(args, ctx)
    path = _edit_target(ctx, args.path)
    original = path.read_text(encoding="utf-8-sig")
    backup_dir = ctx.data_dir / "backups" / ctx.task_id
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = unique_path(backup_dir / path.name)
    shutil.copy2(path, backup)
    count = original.count(args.find) if args.replace_all else 1
    updated = original.replace(args.find, args.replace_with, -1 if args.replace_all else 1)
    tmp = path.with_name(f".{path.name}.claw-tmp")
    tmp.write_text(updated, encoding="utf-8", newline="")
    tmp.replace(path)
    ok = path.read_text(encoding="utf-8") == updated
    _record_output(ctx, path, "edited", ok, f"Backup of the original: {backup}")
    return ToolResult(ok, f"Changed {count} place(s) in {path.name}. The original was backed up first.",
                      {"path": str(path), "backup": str(backup), "changes": count}, verified=ok)


def register_document_tools(reg: ToolRegistry) -> None:
    F, W = Capability.FILES, RiskLevel.LOCAL_WRITE
    reg.register(ToolSpec("write_text_file", "Create a new text file (.txt .md .csv .json .html) in an authorized folder.",
                          F, W, WriteTextArgs, write_text_file, precheck=check_write_text))
    reg.register(ToolSpec("save_report", "Save a report (Word or Markdown) with sections and the sources Claw opened.",
                          F, W, ReportArgs, save_report, precheck=check_report))
    reg.register(ToolSpec("create_spreadsheet", "Create a spreadsheet (Excel or CSV) from rows of data Claw collected.",
                          F, W, SheetArgs, create_spreadsheet, precheck=check_sheet))
    reg.register(ToolSpec("edit_text_file", "Replace text in an existing text file (a backup is kept).",
                          F, W, EditArgs, edit_text_file, precheck=check_edit))
