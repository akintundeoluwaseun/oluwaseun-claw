"""Turn documents into plain text: text files, PDF (pypdf), Word (python-docx), Excel (openpyxl)."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from claw.tools.base import ToolError

TEXT_TYPES = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".log", ".xml", ".html", ".htm", ".css", ".js", ".ts", ".py",
    ".ini", ".cfg", ".yaml", ".yml", ".toml", ".sql", ".srt", ".vtt",
}
DOC_TYPES = {".pdf", ".docx", ".xlsx", ".xlsm"}
READABLE = TEXT_TYPES | DOC_TYPES
MAX_BYTES = 40_000_000
MAX_PDF_PAGES = 200
MAX_SHEET_ROWS = 2000


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ToolError("The file's text encoding could not be read.")


def _pdf(path: Path) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ToolError("This PDF is password-protected.")
            except Exception as exc:  # noqa: BLE001
                raise ToolError("This PDF is password-protected.") from exc
        parts = []
        for i, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(f"[Page {i}]\n{text}")
    except PdfReadError as exc:
        raise ToolError(f"The PDF could not be read: {exc}") from exc
    if not parts:
        raise ToolError("This PDF has no readable text (it may be a scanned image).")
    return "\n\n".join(parts)


def _docx(path: Path) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        doc = docx.Document(str(path))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise ToolError("The Word document could not be opened.") from exc
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for t_index, table in enumerate(doc.tables, 1):
        lines.append(f"[Table {t_index}]")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises several types for bad files
        raise ToolError("The spreadsheet could not be opened.") from exc
    out = []
    try:
        for ws in wb.worksheets:
            out.append(f"[Sheet: {ws.title}]")
            buf = io.StringIO()
            writer = csv.writer(buf)
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_SHEET_ROWS:
                    buf.write(f"... (more rows not shown)\n")
                    break
                if any(v is not None for v in row):
                    writer.writerow(["" if v is None else v for v in row])
            out.append(buf.getvalue().rstrip())
    finally:
        wb.close()
    return "\n".join(out)


def read_document(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in READABLE:
        supported = "text files, PDF, Word (.docx) and Excel (.xlsx)"
        raise ToolError(f"Claw can read {supported}; '{ext or 'no extension'}' isn't supported.")
    size = path.stat().st_size
    if size > MAX_BYTES:
        raise ToolError(f"The file is too large to read ({size:,} bytes).")
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(path)
    return _decode(path.read_bytes())
