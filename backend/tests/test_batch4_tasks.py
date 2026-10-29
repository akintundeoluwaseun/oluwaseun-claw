"""Batch 4 acceptance: multi-step tasks through the real agent loop.
Model replies are scripted (TEST DOUBLE); browser, website, files and checks are real."""

import json
import time

import docx
from openpyxl import load_workbook

from claw.agent.executor import run_agent
from claw.models.ollama import OllamaAdapter
from claw.permissions.policy import Decision, PermissionPolicy, RiskLevel
from claw.tools.registry import build_registry
from tests.site_server import DIRECTORY
from tests.test_documents import make_pdf

REG = build_registry()


def act(tool, **args):
    return json.dumps({"thought": f"use {tool}", "action": "tool", "tool": tool, "args": args})


def finish(text):
    return json.dumps({"thought": "done", "action": "finish", "final_answer": text})


class Rec:
    def __init__(self, answer=True):
        self.events, self.approvals, self.answer = [], [], answer

    def emit(self, k, d):
        self.events.append((k, d))

    async def approve(self, p):
        self.approvals.append(p)
        return self.answer


async def run(ctx, fake_server, replies, task, answer=True):
    fake_server.replies.extend(replies)
    rec = Rec(answer)
    adapter = OllamaAdapter(fake_server.base_url, "qwen2.5:7b", timeout_s=10, temperature=0)
    out = await run_agent(task, "", adapter, REG, ctx, rec.emit, rec.approve)
    return out, rec


async def test_research_and_save_report_with_sources(tool_ctx, fake_server, site):
    url = site.base_url + "/articles/motion-templates"
    url2 = site.base_url + "/articles/gsap-starter"
    out, rec = await run(tool_ctx, fake_server, [
        act("web_search", query="free motion graphics templates"),
        act("open_search_result", number=1),
        act("read_page"),
        act("open_search_result", number=2),
        act("read_page"),
        act("save_report", title="Free motion graphics templates", format="docx",
            sections=[{"heading": "Summary", "body": "Two free packs were found: one for After Effects and HTML5, one GSAP starter."}],
            sources=[{"title": "Free Motion Graphics Templates 2026", "location": url},
                     {"title": "GSAP Starter Pack", "location": url2}]),
        finish("Saved a report with 2 sources."),
    ], "Research free motion graphics templates and save a report with the source links")
    assert out.status == "completed", out.summary()
    assert all(s["outcome"] == "verified" for s in out.steps)
    [f] = out.files
    assert f["kind"] == "report" and f["verified"]
    text = "\n".join(p.text for p in docx.Document(f["path"]).paragraphs)
    assert url in text and url2 in text


async def test_find_files_and_summarize(tool_ctx, fake_server, tmp_path):
    folder = tmp_path / "saved"
    make_pdf(folder / "old.pdf", "Old notes")
    time.sleep(0.05)
    make_pdf(folder / "latest.pdf", "Board approved budget of 2 million USD")
    out, rec = await run(tool_ctx, fake_server, [
        act("find_files", extension=".pdf"),
        act("read_file", path=str(folder / "latest.pdf")),
        act("write_text_file", file_name="latest summary.md", text="# Summary\nThe board approved a budget of 2 million USD."),
        finish("Summarised latest.pdf into latest summary.md."),
    ], "Find the latest PDF in my folder and summarize it")
    assert out.status == "completed" and out.files[0]["path"].endswith("latest summary.md")
    assert out.steps[0]["data"]["files"][0]["path"].endswith("latest.pdf")  # newest first


async def test_extract_webpage_to_spreadsheet_and_injection_guard(tool_ctx, fake_server, site):
    # The user allows file saving without asking; the page then tries to inject instructions.
    tool_ctx.policy = PermissionPolicy({RiskLevel.LOCAL_WRITE: Decision.ALLOW}, tool_ctx.settings.permissions.authorized_folders)
    rows = [list(r) for r in DIRECTORY]
    out, rec = await run(tool_ctx, fake_server, [
        act("open_url", url=site.base_url + "/directory"),
        act("read_page"),
        act("create_spreadsheet", file_name="city businesses", columns=["Name", "Type", "Phone", "Email"], rows=rows),
        finish("Saved 3 businesses."),
    ], "Collect the businesses on this directory page into a spreadsheet")
    assert out.status == "completed", out.summary()
    assert any(k == "injection_warning" for k, _ in rec.events)
    # Saving would normally go ahead without asking, but the page tried to give orders: Claw asked.
    assert [a["tool"] for a in rec.approvals] == ["create_spreadsheet"] and "tried to give it instructions" in rec.approvals[0]["caution"]
    assert out.warnings and "tried to give it instructions" in out.warnings[0]
    ws = load_workbook(out.files[0]["path"]).active
    assert [c.value for c in ws[2]] == list(DIRECTORY[0])
    assert out.files[0]["note"] == ""  # every value was found on the page
    sent = fake_server.requests[-1]["body"]["messages"]
    assert any("WARNING: this content contains text that tries to give you instructions" in m["content"] for m in sent)


async def test_reading_web_content_makes_online_changes_ask(tool_ctx, fake_server, site):
    tool_ctx.policy = PermissionPolicy({RiskLevel.EXTERNAL: Decision.ALLOW}, [])
    out, rec = await run(tool_ctx, fake_server, [
        act("open_url", url=site.base_url + "/articles/gsap-starter"),
        act("read_page"),
        act("open_url", url=site.base_url + "/contact"),
        act("click_button", text="Go"),
    ], "Read the article then use the contact page", answer=False)
    # 'Go' is not a consequential word and external actions are allowed, but web content was read first.
    assert rec.approvals and rec.approvals[0]["tool"] == "click_button"
    assert "changes outside this computer" in rec.approvals[0]["caution"]
    assert out.status == "stopped"


async def test_recovers_by_replanning(tool_ctx, fake_server, site):
    out, rec = await run(tool_ctx, fake_server, [
        act("open_url", url=site.base_url + "/missing-page"),
        act("open_url", url=site.base_url + "/also-missing"),
        json.dumps({"diagnosis": "Guessed addresses do not exist; search instead.", "new_plan": ["Search the site", "Open the best result"]}),
        act("web_search", query="gsap motion graphics"),
        act("open_search_result", number=1),
        finish("Found the GSAP starter pack."),
    ], "Find the GSAP starter pack")
    assert out.status == "completed" and out.replans == 1
    assert [s["outcome"] for s in out.steps] == ["failed", "failed", "verified", "verified"]
    replanned = [d for k, d in rec.events if k == "replanned"][0]
    assert replanned["new_plan"][0] == "Search the site"
