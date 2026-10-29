"""Document reading and writing tools (real files, real libraries)."""

import csv
from pathlib import Path

import docx
import pytest
from openpyxl import Workbook, load_workbook

from claw.documents.pdf_pipeline import PdfPipeline
from claw.documents.readers import read_document
from claw.tools.base import ToolError
from claw.tools.registry import build_registry

REG = build_registry()


async def call(ctx, tool_name, **args):
    spec = REG.get(tool_name)
    parsed = spec.parse_args(args)
    if spec.precheck:
        spec.precheck(parsed, ctx)
    return await spec.handler(parsed, ctx)


def make_pdf(path, text):
    """Minimal one-page PDF with real text (no extra libraries)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))


@pytest.fixture
def folder(tool_ctx, tmp_path):
    return tmp_path / "saved"


def test_read_pdf_docx_xlsx(tmp_path):
    make_pdf(tmp_path / "a.pdf", "Invoice total 45000 USD")
    assert "Invoice total 45000 USD" in read_document(tmp_path / "a.pdf")
    d = docx.Document()
    d.add_paragraph("Meeting notes for Monday")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Item", "Owner"
    d.save(tmp_path / "b.docx")
    text = read_document(tmp_path / "b.docx")
    assert "Meeting notes for Monday" in text and "Item | Owner" in text
    wb = Workbook()
    wb.active.append(["Name", "Phone"])
    wb.active.append(["Ade", "0803"])
    wb.save(tmp_path / "c.xlsx")
    assert "Ade,0803" in read_document(tmp_path / "c.xlsx")
    (tmp_path / "d.exe").write_bytes(b"MZ")
    with pytest.raises(ToolError, match="isn't supported"):
        read_document(tmp_path / "d.exe")


async def test_read_file_tool_records_seen(tool_ctx, folder):
    make_pdf(folder / "report.pdf", "Quarterly sales grew 12 percent")
    r = await call(tool_ctx, "read_file", path=str(folder / "report.pdf"))
    # saw_file, not `in seen_files`: the recorded path is canonical, so on Windows the set
    # holds a case-folded path and a raw string match would miss the file Claw just read.
    assert "12 percent" in r.untrusted_text
    assert tool_ctx.memory.saw_file(str(folder / "report.pdf"))


async def test_write_text_file_never_overwrites(tool_ctx, folder):
    a = await call(tool_ctx, "write_text_file", file_name="summary.md", text="# Hi")
    b = await call(tool_ctx, "write_text_file", file_name="summary.md", text="# Again")
    assert a.verified and b.data["path"].endswith("summary (2).md")
    assert (folder / "summary.md").read_text() == "# Hi"
    with pytest.raises(ToolError):
        REG.get("write_text_file").parse_args({"file_name": "run.bat", "text": "x"})


async def test_report_only_cites_opened_sources(tool_ctx, folder, site):
    url = site.base_url + "/articles/motion-templates"
    args = dict(title="Motion templates", sections=[{"heading": "Findings", "body": "Two good free packs."}],
                sources=[{"title": "Free templates", "location": url}])
    with pytest.raises(ToolError, match="not opened"):
        await call(tool_ctx, "save_report", **args)
    await call(tool_ctx, "open_url", url=url)
    r = await call(tool_ctx, "save_report", **args)
    assert r.verified and r.data["path"].endswith("Motion templates.docx")
    doc = docx.Document(r.data["path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Findings" in text and url in text and "Sources" in text
    md = await call(tool_ctx, "save_report", **{**args, "format": "md"})
    assert url in open(md.data["path"], encoding="utf-8").read()
    assert [o["kind"] for o in tool_ctx.memory.outputs] == ["report", "report"]


async def test_spreadsheet_provenance_and_formula_safety(tool_ctx, folder):
    tool_ctx.memory.seen_text.append("Riverside Hotel  0803 111 2233  info@riversidegrand.example")
    r = await call(tool_ctx, "create_spreadsheet", file_name="hotels", columns=["Name", "Phone", "Email"],
                   rows=[["Riverside Hotel", "0803 111 2233", "info@riversidegrand.example"],
                         ["Made Up Hotel", "0800 000 0000", "=HYPERLINK(\"http://evil\")"]])
    assert r.verified and r.data["rows"] == 2
    assert r.data["checked"] == 6 and r.data["found"] == 3
    note = tool_ctx.memory.outputs[-1]["note"]
    assert "3 of 6" in note and "Made Up Hotel" in note
    ws = load_workbook(r.data["path"]).active
    assert ws["A1"].value == "Name" and ws["A1"].font.bold
    assert ws["C3"].value.startswith("'=")  # stored as text, not a formula
    c = await call(tool_ctx, "create_spreadsheet", file_name="hotels", format="csv", columns=["Name"], rows=[["Riverside Hotel"]])
    with open(c.data["path"], encoding="utf-8-sig", newline="") as fh:
        assert list(csv.reader(fh)) == [["Name"], ["Riverside Hotel"]]


def test_spreadsheet_row_width_checked():
    with pytest.raises(ToolError, match="row 1 has 3 values"):
        REG.get("create_spreadsheet").parse_args({"file_name": "x", "columns": ["a", "b"], "rows": [[1, 2, 3]]})


async def test_edit_keeps_backup(tool_ctx, folder):
    f = folder / "todo.txt"
    f.write_text("buy rice\nbuy beans\n", encoding="utf-8")
    with pytest.raises(ToolError, match="appears 2 times"):
        await call(tool_ctx, "edit_text_file", path=str(f), find="buy", replace_with="get")
    with pytest.raises(ToolError, match="not found"):
        await call(tool_ctx, "edit_text_file", path=str(f), find="yam", replace_with="x")
    r = await call(tool_ctx, "edit_text_file", path=str(f), find="buy rice", replace_with="buy ofada rice")
    assert r.verified and f.read_text() == "buy ofada rice\nbuy beans\n"
    assert open(r.data["backup"]).read() == "buy rice\nbuy beans\n"


# --- Regression: a source Claw really opened must be citable however the path is spelled. ---
# The recorded path is canonical (policy.canonical_path): '~', '..' and symlinks are resolved,
# and on Windows the case is folded. Comparing the model's raw spelling against that set made
# save_report reject a file it had just read -- on Windows for every path with a capital letter
# (test_read_file_tool_records_seen failed there), and on any platform for a '..' path.

async def test_seen_file_is_citable_through_a_dotdot_path(tool_ctx, folder):
    (folder / "sub").mkdir()
    make_pdf(folder / "report.pdf", "Quarterly sales grew 12 percent")
    spelled = str(folder / "sub" / ".." / "report.pdf")
    await call(tool_ctx, "read_file", path=spelled)
    args = dict(title="Quarterly", sections=[{"heading": "Findings", "body": "Sales grew."}],
                sources=[{"title": "Report", "location": spelled}])
    r = await call(tool_ctx, "save_report", **args)
    assert r.verified


def test_memory_matches_paths_by_identity_not_spelling(tmp_path):
    from claw.tools.base import TaskMemory

    f = tmp_path / "a.pdf"
    f.write_bytes(b"%PDF-1.4\n")
    m = TaskMemory()
    m.note_file(f)
    assert m.saw_file(str(f))
    assert m.saw_file(str(tmp_path / "sub" / ".." / "a.pdf"))
    assert not m.saw_file(str(tmp_path / "b.pdf"))
    assert not m.saw_file("https://example.com/a.pdf")  # a URL is not a path
    assert not m.saw_file("")


def test_memory_folds_case_where_the_filesystem_does(tmp_path, monkeypatch):
    """Emulates Windows: os.path.normcase folds case, so the two spellings are one file."""
    import ntpath

    from claw.tools.base import TaskMemory

    monkeypatch.setattr("claw.permissions.policy.os.path.normcase", ntpath.normcase)
    m = TaskMemory()
    m.note_file(tmp_path / "Report.PDF")
    assert m.saw_file(str(tmp_path / "report.pdf"))


async def test_unseen_source_is_still_refused(tool_ctx, folder):
    make_pdf(folder / "report.pdf", "Quarterly sales grew 12 percent")
    with pytest.raises(ToolError, match="not opened"):
        await call(tool_ctx, "save_report", title="Quarterly",
                   sections=[{"heading": "Findings", "body": "Sales grew."}],
                   sources=[{"title": "Report", "location": str(folder / "report.pdf")}])


# --- Extraction content: whitespace, newlines, Unicode and numbers survive the read. ---

@pytest.mark.parametrize("label,text,needles", [
    ("plain", "Quarterly sales grew 12 percent", ["12 percent"]),
    ("percent_sign", "Operating margin rose to 12.5% this financial year", ["12.5%"]),
    ("money_and_date", "Total due: 45,000 USD by 30 November 2026", ["45,000", "30 November 2026"]),
    ("hyphen_and_name", "Riverside branch grew 7 percent", ["Riverside", "7 percent"]),
])
def test_text_pdf_keeps_numbers_and_characters(tmp_path, label, text, needles):
    make_pdf(tmp_path / f"{label}.pdf", text)
    r = PdfPipeline(tmp_path / "cache").read(tmp_path / f"{label}.pdf")
    got = r.pages[0].text
    assert r.pages[0].source == "text"
    for n in needles:
        assert n in got, f"{label}: {n!r} missing from {got!r}"


def test_crlf_and_lf_normalise_to_the_same_text(tmp_path):
    out = {}
    for name, nl in (("crlf", "\r\n"), ("lf", "\n")):
        p = tmp_path / f"{name}.txt"
        p.write_bytes(f"Quarterly sales{nl}grew 12 percent{nl}".encode())
        out[name] = read_document(p).strip()
    assert out["crlf"] == out["lf"]
    assert "\r" not in out["crlf"]


def test_empty_and_malformed_documents_are_reported_not_silently_empty(tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"%PDF-1.4\nnot really a pdf\n")
    with pytest.raises(ToolError):
        PdfPipeline(tmp_path / "cache").read(bad)


def test_text_reader_keeps_unicode_and_non_breaking_spaces(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text("Caf\u00e9 Riverside\u2014total 45\u00a0000 USD\n", encoding="utf-8")
    got = read_document(p)
    assert "Caf\u00e9" in got and "Riverside" in got
    assert "45\u00a0000" in got or "45 000" in got  # NBSP kept, or folded to a plain space


@pytest.mark.parametrize("name,expect", [
    ("contacts.csv", ".csv"),     # the name says csv, so it is a csv
    ("contacts.xlsx", ".xlsx"),
    ("contacts", ".xlsx"),        # no extension: the format default applies
])
async def test_spreadsheet_name_decides_the_kind_of_file(tool_ctx, folder, name, expect):
    """'contacts.csv' used to produce 'contacts.csv.xlsx' - a binary file with two extensions."""
    r = await call(tool_ctx, "create_spreadsheet", file_name=name, columns=["A", "B"],
                   rows=[["1", "2"]], folder=str(folder))
    out = Path(r.data["path"])
    assert out.suffix == expect, out.name
    assert out.name.count(".") == 1, f"double extension: {out.name}"
    if expect == ".csv":
        assert out.read_text(encoding="utf-8-sig").splitlines()[0] == "A,B"
    else:
        assert out.read_bytes()[:2] == b"PK"  # a real xlsx is a zip
